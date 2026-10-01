"""
The workbook the FIDE registration form fills in (kbsb.fide.api_fide.fill_workbook).

Its values are typed on a public form, and fide@ and the staff open the file
in Excel, so a value that looks like a formula must be stored as text (SEC-40),
without changing the text: the dataplatform reads these cells. The template's
own formulas, in cells the form does not fill, stay formulas.
"""

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from kbsb.core.cells import write_text
from kbsb.fide.api_fide import FIDE_FIELDS, fill_workbook

MALICIOUS = [
    "=1+1",
    '=HYPERLINK("http://x")',
    "+cmd|' /C calc'!A0",
    "@SUM(1)",
    "-2+3",
    "=cmd|' /C calc'!A0",
]

FIRST_ROW = 10
ROW = {key: FIRST_ROW + i for i, (key, _) in enumerate(FIDE_FIELDS)}

# the template's own formulas next to the FIDE-ID cells (the licence check)
TEMPLATE_FORMULAS = ["C28", "C30", "C32", "C35", "C37", "C39", "C41"]


def reload(buf):
    # not data_only: a formula would come back as its formula text
    return load_workbook(BytesIO(buf.getvalue()))


def malicious_form():
    """Every field of the form filled with one of the payloads."""
    form = {
        "invoice_email": "=1+1",
        "invoice_clubnr": '=HYPERLINK("http://x")',
    }
    for i, (key, _) in enumerate(FIDE_FIELDS):
        form[key] = MALICIOUS[i % len(MALICIOUS)]
    return form


def test_formulas_from_the_form_are_stored_as_text():
    form = malicious_form()
    ws = reload(fill_workbook(form))["FIDE Registration Form"]

    for ref, key in (("B7", "invoice_email"), ("B8", "invoice_clubnr")):
        assert ws[ref].data_type == "s", ref
        assert ws[ref].value == form[key]
    for key, _ in FIDE_FIELDS:
        if key == "multiple_round_days":
            continue  # always left empty
        cell = ws[f"B{ROW[key]}"]
        assert cell.data_type == "s", key
        assert cell.value == form[key], key


@pytest.mark.parametrize("payload", MALICIOUS)
def test_each_payload_in_the_event_name(payload):
    ws = reload(fill_workbook({"event_name": payload}))["FIDE Registration Form"]
    cell = ws[f"B{ROW['event_name']}"]
    assert cell.data_type == "s"
    assert cell.value == payload


def test_leading_tab_and_line_breaks_are_dropped():
    form = {"event_name": "\t=1+1", "city": "\r\n@SUM(1)", "remarks": "a\tb"}
    ws = reload(fill_workbook(form))["FIDE Registration Form"]
    assert ws[f"B{ROW['event_name']}"].value == "=1+1"
    assert ws[f"B{ROW['event_name']}"].data_type == "s"
    assert ws[f"B{ROW['city']}"].value == "@SUM(1)"
    # only a leading one goes
    assert ws[f"B{ROW['remarks']}"].value == "a\tb"


def test_normal_values_are_unchanged():
    form = {
        "invoice_email": "club@example.be",
        "invoice_clubnr": "601",
        "event_name": "Brussels Open 2027",
        "city": "Sint-Niklaas",
        "start_date": "2027-03-10",
        "expected_players": 120,
        "remarks": "+32 478 12 34 56",
        "communication_language": "fr",
    }
    ws = reload(fill_workbook(form))["FIDE Registration Form"]
    assert ws["B7"].value == "club@example.be"
    assert ws["B8"].value == "601"
    assert ws[f"B{ROW['event_name']}"].value == "Brussels Open 2027"
    assert ws[f"B{ROW['city']}"].value == "Sint-Niklaas"
    assert ws[f"B{ROW['start_date']}"].value == "2027-03-10"
    # a number stays a number
    assert ws[f"B{ROW['expected_players']}"].value == 120
    assert ws[f"B{ROW['expected_players']}"].data_type == "n"
    assert ws[f"B{ROW['remarks']}"].value == "+32 478 12 34 56"
    assert ws[f"B{ROW['communication_language']}"].value == "Français (FR)"


def test_template_formulas_stay_formulas():
    wb = reload(fill_workbook(malicious_form()))
    ws = wb["FIDE Registration Form"]
    for ref in TEMPLATE_FORMULAS:
        assert ws[ref].data_type == "f", ref
        assert str(ws[ref].value).startswith("=IF("), ref
    # not a long tournament: the rounds sheet keeps its own formula
    assert wb["Rounds_Long_Tournament"]["B1"].data_type == "f"


def test_rounds_sheet_stores_text():
    form = {
        "tournament_report": "New long tournament",
        "event_name": '=HYPERLINK("http://x")',
        "rounds_reported": "2",
        "round1_date": "=1+1",
        "round1_report": "@SUM(1)",
        "round1_end_date": "-2+3",
        "round2_date": "2027-03-17",
        "round2_report": "1",
    }
    ws = reload(fill_workbook(form))["Rounds_Long_Tournament"]
    expected = {
        "B1": '=HYPERLINK("http://x")',
        "D1": "End Date (optional)",
        "A2": "Round 1 Date",
        "B2": "=1+1",
        "C2": "@SUM(1)",
        "D2": "-2+3",
        "B3": "2027-03-17",
        "C3": "1",
    }
    for ref, value in expected.items():
        assert ws[ref].data_type == "s", ref
        assert ws[ref].value == value, ref
    assert ws["D3"].value is None


def test_write_text_types():
    ws = Workbook().active
    assert write_text(ws["A1"], "=1+1").data_type == "s"
    assert write_text(ws["A2"], 7).value == 7
    assert write_text(ws["A3"], None).value is None
    assert write_text(ws["A4"], ["=1", "x"]).value == "['=1', 'x']"
    assert ws["A4"].data_type == "s"
